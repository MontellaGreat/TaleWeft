# Modified for TaleWeft on 2026-10-22: product branding; compatibility identifiers retained.
# Modified for TaleWeft on 2026-10-22: MIT EPUB writer, complete navigation and XML-safe content.
"""导出服务：生成真实可打开的 DOCX / EPUB / PDF / Markdown。"""
from __future__ import annotations

import html
import io
import logging
import re
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from typing import List, Tuple

from domain.novel.repositories.novel_repository import NovelRepository
from domain.novel.repositories.chapter_repository import ChapterRepository
from domain.novel.entities.novel import Novel
from domain.novel.entities.chapter import Chapter
from domain.novel.value_objects.novel_id import NovelId
from domain.novel.value_objects.chapter_id import ChapterId
from infrastructure.export.font_environment import ExportFontEnvironmentSettings

logger = logging.getLogger(__name__)


def _safe_filename_stem(title: str, max_len: int = 80) -> str:
    t = (title or "novel").strip()
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", t)
    t = t.replace(" ", "_").strip("._") or "novel"
    if len(t) > max_len:
        t = t[:max_len]
    return t


def _novel_id_str(novel: Novel) -> str:
    nid = novel.id
    return nid.value if hasattr(nid, "value") else str(nid)


def _chapter_display_title(ch: Chapter) -> str:
    if ch.title and str(ch.title).strip():
        return str(ch.title).strip()
    return f"第 {ch.number} 章"


def _content_to_html_paragraphs(text: str) -> str:
    raw = _xml_safe_text(text).replace("\r\n", "\n").replace("\r", "\n")
    parts: List[str] = []
    for block in raw.split("\n"):
        line = block.strip()
        if line:
            parts.append(f"<p>{html.escape(line)}</p>")
    if not parts:
        return "<p></p>"
    return "\n".join(parts)


def _xml_safe_text(text: str) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "\ufffd", text or "")


class ExportService:
    """导出服务"""

    def __init__(
        self,
        novel_repository: NovelRepository,
        chapter_repository: ChapterRepository,
        *,
        font_settings: ExportFontEnvironmentSettings | None = None,
    ):
        self.novel_repository = novel_repository
        self.chapter_repository = chapter_repository
        self.font_settings = font_settings or ExportFontEnvironmentSettings.from_env()

    def export_novel(self, novel_id: str, format: str) -> Tuple[bytes, str, str]:
        try:
            logger.info("开始导出小说: %s, 格式: %s", novel_id, format)
            novel = self.novel_repository.get_by_id(NovelId(novel_id))
            if not novel:
                raise ValueError(f"小说不存在: {novel_id}")
            chapters = self.chapter_repository.list_by_novel(NovelId(novel_id))
            chapters.sort(key=lambda x: x.number)
            logger.info("导出: %s, 章节数 %s", novel.title, len(chapters))
            if format == "epub":
                result = self._export_to_epub(novel, chapters)
            elif format == "pdf":
                result = self._export_to_pdf(novel, chapters)
            elif format == "docx":
                result = self._export_to_docx(novel, chapters)
            elif format == "markdown":
                result = self._export_to_markdown(novel, chapters)
            elif format == "json":
                result = self._export_to_json(novel, chapters)
            else:
                raise ValueError(f"不支持的导出格式: {format}")
            logger.info("导出成功，%s 字节", len(result[0]))
            return result
        except ValueError:
            raise
        except Exception as e:
            logger.error("导出小说失败: %s", e, exc_info=True)
            raise

    def export_chapter(self, chapter_id: str, format: str) -> Tuple[bytes, str, str]:
        try:
            logger.info("开始导出章节: %s, 格式: %s", chapter_id, format)
            chapter = self.chapter_repository.get_by_id(ChapterId(chapter_id))
            if not chapter:
                raise ValueError(f"章节不存在: {chapter_id}")
            novel_id = chapter.novel_id.value if hasattr(chapter.novel_id, "value") else chapter.novel_id
            novel = self.novel_repository.get_by_id(NovelId(novel_id))
            if not novel:
                raise ValueError(f"小说不存在: {novel_id}")
            if format == "epub":
                result = self._export_to_epub(novel, [chapter])
            elif format == "pdf":
                result = self._export_to_pdf(novel, [chapter])
            elif format == "docx":
                result = self._export_to_docx(novel, [chapter])
            elif format == "markdown":
                result = self._export_to_markdown(novel, [chapter])
            elif format == "json":
                result = self._export_to_json(novel, [chapter])
            else:
                raise ValueError(f"不支持的导出格式: {format}")
            data, mime, _ = result
            ext = {"epub": "epub", "pdf": "pdf", "docx": "docx", "markdown": "md", "json": "json"}[format]
            chapter_stem = _safe_filename_stem(
                f"{novel.title or 'novel'}-第{chapter.number}章"
            )
            logger.info("导出成功，%s 字节", len(data))
            return data, mime, f"{chapter_stem}.{ext}"
        except ValueError:
            raise
        except Exception as e:
            logger.error("导出章节失败: %s", e, exc_info=True)
            raise

    def _export_to_json(self, novel, chapters):
        from application.core.services.manuscript_transfer import make_document, validate_document
        import json
        text = json.dumps(make_document(novel, chapters), ensure_ascii=False, indent=2)
        validate_document(text)
        return text.encode('utf-8'), 'application/json; charset=utf-8', f'{_safe_filename_stem(novel.title)}.taleweft.json'

    def _export_to_epub(self, novel: Novel, chapters: list[Chapter]) -> Tuple[bytes, str, str]:
        from epublib import EPUB
        from epublib.create import EPUBCreator
        from epublib.nav.util import TOCEntryData
        from epublib.resources.create import create_resource

        def content_document(title: str, body: str) -> bytes:
            return f"""<?xml version="1.0" encoding="utf-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" lang="zh" xml:lang="zh">
<head><title>{html.escape(_xml_safe_text(title))}</title><meta charset="utf-8"/></head>
<body>{body}</body>
</html>""".encode("utf-8")

        buf = io.BytesIO()
        # Keep publication resources at one level: archive URIs stay portable on Windows.
        template = EPUBCreator(navigation_document_path="nav.xhtml").to_file()
        with EPUB(template) as book:
            book.metadata.identifier = _xml_safe_text(f"taleweft:{_novel_id_str(novel)}")
            book.metadata.title = _xml_safe_text(novel.title or "未命名")
            book.metadata.language = "zh"
            book.metadata.add_dc("creator", _xml_safe_text(novel.author or "未知作者"))
            book.metadata.modified = datetime.now(timezone.utc)
            book.nav.soup.html["lang"] = "zh"
            book.nav.soup.html["xml:lang"] = "zh"
            book.nav.soup.title.string = "目录"
            # The navigation remains accessible, while reading starts at the introduction.
            book.spine.items[0].tag["linear"] = "no"

            title = html.escape(_xml_safe_text(novel.title or "未命名"))
            author = html.escape(_xml_safe_text(novel.author or "—"))
            premise = _content_to_html_paragraphs((novel.premise or "").strip() or "（无简介）")
            # A non-linear spine document must be reachable from reading content.
            intro = content_document(
                "简介", f'<h1>{title}</h1><nav><a href="nav.xhtml">目录</a></nav>'
                f"<p>作者：{author}</p>{premise}",
            )
            book.resources.add(create_resource(intro, "intro.xhtml"), add_to_toc=False)
            toc = [TOCEntryData(filename="intro.xhtml", label="简介")]

            for index, chapter in enumerate(chapters, 1):
                chapter_title = _xml_safe_text(_chapter_display_title(chapter))
                body = f"<h1>{html.escape(chapter_title)}</h1>" + _content_to_html_paragraphs(chapter.content)
                document = content_document(chapter_title, body)
                filename = f"chap_{index:03d}.xhtml"
                book.resources.add(create_resource(document, filename), add_to_toc=False)
                toc.append(TOCEntryData(filename=filename, label=chapter_title))

            book.nav.reset_toc(toc)
            book.nav.toc.text = "目录"
            # Build the optional EPUB 2 navigation with XML APIs to escape user titles.
            ncx = ET.Element("ncx", {
                "xmlns": "http://www.daisy.org/z3986/2005/ncx/",
                "version": "2005-1", "xml:lang": "zh",
            })
            head = ET.SubElement(ncx, "head")
            for name, value in (("dtb:uid", book.metadata.identifier), ("dtb:depth", "1"),
                                ("dtb:totalPageCount", "0"), ("dtb:maxPageNumber", "0")):
                ET.SubElement(head, "meta", {"name": name, "content": value})
            ET.SubElement(ET.SubElement(ncx, "docTitle"), "text").text = book.metadata.title
            ET.SubElement(ET.SubElement(ncx, "docAuthor"), "text").text = _xml_safe_text(novel.author or "未知作者")
            nav_map = ET.SubElement(ncx, "navMap")
            for index, entry in enumerate(toc, 1):
                point = ET.SubElement(nav_map, "navPoint", {"id": f"nav_{index}", "playOrder": str(index)})
                ET.SubElement(ET.SubElement(point, "navLabel"), "text").text = entry.label
                ET.SubElement(point, "content", {"src": entry.filename})
            book.resources.add(create_resource(ET.tostring(ncx, encoding="utf-8", xml_declaration=True), "toc.ncx"),
                               add_to_spine=False, add_to_toc=False)
            book.spine.tag["toc"] = book.manifest["toc.ncx"].id
            book.write(buf)

        stem = _safe_filename_stem(novel.title)
        return buf.getvalue(), "application/epub+zip", f"{stem}.epub"

    def _try_register_cjk_font(self, pdf) -> bool:
        for path in self.font_settings.cjk_font_paths():
            if not path.is_file():
                continue
            try:
                pdf.add_font("PlotExportCJK", "", str(path), uni=True)
                return True
            except Exception as e:
                logger.debug("PDF 跳过字体 %s: %s", path, e)
        return False

    def _export_to_pdf(self, novel: Novel, chapters: list[Chapter]) -> Tuple[bytes, str, str]:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=14)
        font = "Helvetica"
        if self._try_register_cjk_font(pdf):
            font = "PlotExportCJK"

        def add_text(size: float, text: str, line_h: float) -> None:
            pdf.set_font(font, size=size)
            body = (text or "").strip() or " "
            try:
                pdf.multi_cell(0, line_h, body, new_x="LMARGIN", new_y="NEXT")
            except Exception as e:
                logger.warning("PDF multi_cell 回退: %s", e)
                pdf.set_font("Helvetica", size=size)
                safe = (text or "").encode("ascii", errors="replace").decode("ascii")
                pdf.multi_cell(0, line_h, safe or " ", new_x="LMARGIN", new_y="NEXT")

        pdf.add_page()
        add_text(16, novel.title or "未命名", 9)
        pdf.ln(2)
        add_text(
            11,
            f"作者：{novel.author or '—'}\n简介：{(novel.premise or '').strip() or '—'}",
            6,
        )
        pdf.ln(4)

        for ch in chapters:
            add_text(14, _chapter_display_title(ch), 8)
            pdf.ln(1)
            add_text(11, (ch.content or "").strip() or "（无正文）", 6)
            pdf.ln(6)

        out = pdf.output()
        if isinstance(out, str):
            data = out.encode("latin-1")
        elif isinstance(out, bytearray):
            data = bytes(out)
        else:
            data = out
        stem = _safe_filename_stem(novel.title)
        return data, "application/pdf", f"{stem}.pdf"

    def _export_to_docx(self, novel: Novel, chapters: list[Chapter]) -> Tuple[bytes, str, str]:
        from docx import Document

        doc = Document()
        doc.add_heading(novel.title or "未命名", level=0)
        doc.add_paragraph(f"作者：{novel.author or '—'}")
        p_pre = doc.add_paragraph()
        p_pre.add_run("简介：").bold = True
        p_pre.add_run((novel.premise or "").strip() or "（无）")

        for ch in chapters:
            doc.add_heading(_chapter_display_title(ch), level=1)
            content = ch.content or ""
            if not content.strip():
                doc.add_paragraph("（无正文）")
                continue
            for line in content.splitlines():
                doc.add_paragraph(line)

        buf = io.BytesIO()
        doc.save(buf)
        stem = _safe_filename_stem(novel.title)
        return (
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{stem}.docx",
        )

    def _export_to_markdown(self, novel: Novel, chapters: list[Chapter]) -> Tuple[bytes, str, str]:
        lines: List[str] = [
            f"# {novel.title or '未命名'}",
            "",
            f"**作者**: {novel.author or '—'}",
            "",
            "## 简介",
            "",
            (novel.premise or "").strip() or "（无）",
            "",
        ]
        for ch in chapters:
            lines.append(f"## {_chapter_display_title(ch)}")
            lines.append("")
            lines.append((ch.content or "").strip() or "（无正文）")
            lines.append("")
        text = "\n".join(lines)
        stem = _safe_filename_stem(novel.title)
        return text.encode("utf-8"), "text/markdown; charset=utf-8", f"{stem}.md"
